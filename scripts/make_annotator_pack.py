"""Make a blank annotator pack (Excel with instructions + y/n drop-down) from an existing claim-support sheet.
  python scripts/make_annotator_pack.py --sheet data/review/faith1_A.csv --out data/review/faith1_annotator_D.xlsx
Only the id, claim and passages columns are copied; every decision and note is blanked."""
import argparse, csv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TEXT = ["Claim-support labelling: instructions", "",
"You will judge 200 short claims taken from answers to technical-support questions (IBM product documentation). Each row has a claim and the passage text the claim refers to. You do not know which system wrote which claim.", "",
"QUESTION for every row: is the claim fully supported by the passages shown in the same row?",
"TASK: in the sheet 'Items', fill the column 'decision (y/n)' with y or n for every row (drop-down). 'notes' is optional.", "",
"Write y only if everything the claim says can be found in, or follows directly from, the passages shown. A reader with only those passages must be able to confirm the claim without outside knowledge.",
"Write n if: (1) the passages do not mention what the claim says, even if it is true in general; (2) the claim adds a step, value, setting, tool, version or warning that the passages do not contain; (3) the claim contradicts the passages or is about a different product or version; (4) the claim is only partly supported (one part is in the passages and another is not).", "",
"EASY MISTAKES: do not use your own knowledge (if you know it is correct but the passages do not say it, write n). A reasonable paraphrase of the same fact is y, and simple inferences (e.g. 'restart the service' from 'stop and start the service') are y. Generic filler ('Follow the steps below', 'Contact support if the problem continues') is y only if the passages say the same. Each passage is cut at 150 words: judge only what is shown. Labels such as DOCUMENTED or INFERENCE may appear in a claim; ignore them. Do not judge whether the answer is correct or helpful, only whether THIS claim is supported by THESE passages. If unsure after one careful re-read, write n and say why in 'notes'.", "",
"EXAMPLES (invented). Claim 'Set the timeout to 60 seconds.' Passage 'The default timeout is 30 seconds. Increase it in the settings file.' -> n (60 is not in the text).",
"Claim 'Increase the timeout in the settings file.' Same passage -> y.   Claim 'Restart the server and clear the cache.' Passage 'Restart the server.' -> n (clearing the cache is added).   Claim 'The error occurs on version 8.1.' Passage 'This error occurs on version 8.0 only.' -> n.", "",
"PROCESS: work alone and do not discuss rows with anyone until you have returned the file. Expect about two hours. Do not edit the id, claim or passages columns. Save and send the file back as Excel. Thank you."]
ap = argparse.ArgumentParser(); ap.add_argument("--sheet", required=True); ap.add_argument("--out", required=True)
a = ap.parse_args()
rows = list(csv.reader(open(a.sheet, newline="", encoding="utf-8-sig")))
hdr, data = rows[0], rows[1:]
wb = Workbook(); ws = wb.active; ws.title = "Instructions"
for i, t in enumerate(TEXT, 1):
    ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
ws["A1"].font = Font(bold=True, size=13); ws.column_dimensions["A"].width = 120
w2 = wb.create_sheet("Items"); w2.append(hdr[:3] + ["decision (y/n)", "notes"])
for c in w2[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDDDDD")
for r in data: w2.append(r[:3] + ["", ""])
for col, wd in zip("ABCDE", (6, 50, 90, 14, 30)): w2.column_dimensions[col].width = wd
for row in w2.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
dv = DataValidation(type="list", formula1='"y,n"', allow_blank=True); w2.add_data_validation(dv); dv.add(f"D2:D{len(data) + 1}")
w2.freeze_panes = "A2"; wb.save(a.out); print("wrote", a.out, len(data), "claims")
